import json, re
from xml.sax.saxutils import escape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from .core import ROOT


def make_report(name, payload):
    name = re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip(".") or "report"
    directory = ROOT / "storage" / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    pdf = directory / (name + ".pdf")
    docx = directory / (name + ".docx")
    styles = getSampleStyleSheet()
    flow = [Paragraph(escape(name.replace("-", " ")), styles["Title"])]
    document = Document()
    document.add_heading(name.replace("-", " "), 0)
    for key, value in payload.items():
        text = json.dumps(value, indent=2, default=str) if isinstance(value, (dict, list)) else str(value)
        document.add_heading(key.replace("_", " ").title(), 1)
        document.add_paragraph(text)
        flow.extend([Paragraph(escape(key.replace("_", " ").title()), styles["Heading2"])])
        for line in text.splitlines():
            flow.append(Paragraph(escape(line), styles["BodyText"]))
        flow.append(Spacer(1, 10))
    SimpleDocTemplate(str(pdf)).build(flow)
    document.save(docx)
    return {"pdf": pdf.name, "docx": docx.name}
